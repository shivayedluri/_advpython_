from dataclasses import dataclass
from openpyxl import Workbook
"""from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference
wb=Workbook()
sheet=wb.active
rows=[['product','online','store'],
      [20,30,40],
      [30,50,60],
      [20,50,60],
      [50,50,60],
      ]
for row in rows:
    sheet.append(row)
chart=BarChart()
data=Reference(worksheet=sheet,min_row=1,max_row=8,min_col=2,max_col=3)
chart.add_data(data,titles_from_data=True)
sheet.add_chart(chart,"E2")
wb.save('../data/barchat1.xlsx')"""
@dataclass
class people():
      name:str
      no:int
      age:int
obj=[people('string',12,24),people('raju',2,34),people('kamal',3,24)]
data=[[p.name,p.no,p.age]for p in obj]
wb=Workbook()
sheet=wb.active
for i in data:
      sheet.append(i)
wb.save('../data/dataclasses.xlsx')

print(obj)